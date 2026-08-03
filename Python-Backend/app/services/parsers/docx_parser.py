import zipfile
from typing import List

from app.services.parsers.base_parser import BaseParser
from app.services.parsers.parser_factory import ParserFactory


@ParserFactory.register_parser([".docx"])
class DocxParser(BaseParser):
    def _extract_blocks(self, file_path: str, file_name: str) -> List[dict]:
        text = ""
        try:
            import docx

            doc = docx.Document(file_path)
            text = "\n".join([p.text for p in doc.paragraphs])
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " | ".join([cell.text.strip() for cell in row.cells])
        except Exception:
            try:
                import xml.etree.ElementTree as ET

                with zipfile.ZipFile(file_path) as zf:
                    xml_content = zf.read("word/document.xml")
                    root = ET.fromstring(xml_content)
                    text_parts = []
                    for elem in root.iter():
                        if elem.tag.endswith("t"):
                            text_parts.append(elem.text or "")
                    text = "".join(text_parts)
            except Exception as e:
                raise ValueError(f"Failed to extract text from DOCX file: {e}")

        blocks = []
        if text.strip():
            blocks.append(
                {
                    "text": text,
                    "metadata": {
                        "page": 1,
                        "type": "text",
                        "source": "docx_extraction",
                        "char_count": len(text),
                    },
                }
            )

        # Extract images from zip
        try:
            with zipfile.ZipFile(file_path) as zf:
                media_files = [n for n in zf.namelist() if n.startswith("word/media/")]
                for idx, media_file in enumerate(media_files):
                    img_bytes = zf.read(media_file)

                    if len(img_bytes) < 5_000:
                        continue

                    if self.vision_callback:
                        desc = self.vision_callback(
                            img_bytes, f"Image {idx + 1} from DOCX document"
                        )
                    else:
                        desc = "[Image present - vision model disabled]"

                    full_desc = f"[DOCX Image {idx + 1}]:\n{desc}"
                    blocks.append(
                        {
                            "text": full_desc,
                            "metadata": {
                                "page": 1,
                                "type": "image",
                                "source": "docx_image",
                                "image_index": idx,
                            },
                        }
                    )
        except Exception as e:
            print(f"[WARNING] Failed to extract images from DOCX: {e}")

        return blocks
