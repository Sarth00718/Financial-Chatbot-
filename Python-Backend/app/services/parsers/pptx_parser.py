from typing import List

from app.services.parsers.base_parser import BaseParser
from app.services.parsers.parser_factory import ParserFactory


@ParserFactory.register_parser([".pptx"])
class PptxParser(BaseParser):
    def _extract_all_shapes(self, shapes_container):
        """Recursively unnest group shapes to extract all child shapes."""
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        flat_shapes = []
        for shape in shapes_container:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                try:
                    flat_shapes.extend(self._extract_all_shapes(shape.shapes))
                except Exception:
                    pass
            else:
                flat_shapes.append(shape)
        return flat_shapes

    def _extract_blocks(self, file_path: str, file_name: str) -> List[dict]:
        try:
            from pptx import Presentation
            from pptx.enum.shapes import MSO_SHAPE_TYPE
        except ImportError:
            raise ImportError("python-pptx is not installed. Cannot parse PPTX files.")

        prs = Presentation(file_path)
        blocks = []

        for slide_num, slide in enumerate(prs.slides, start=1):
            slide_text = []
            image_count = 0

            shapes = self._extract_all_shapes(slide.shapes)

            for shape in shapes:
                if hasattr(shape, "text") and shape.text and shape.text.strip():
                    slide_text.append(shape.text.strip())

                # Extract image if picture shape or has image attribute
                is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image")
                if is_picture:
                    image_count += 1
                    try:
                        img_bytes = shape.image.blob
                        if len(img_bytes) >= 3_000:
                            if self.vision_callback:
                                desc = self.vision_callback(
                                    img_bytes,
                                    f"Slide {slide_num}, image {image_count} from PPTX presentation",
                                )
                            else:
                                desc = "[Image present - vision model disabled]"

                            full_desc = f"[PPTX Slide {slide_num} - Image {image_count}]:\n{desc}"
                            blocks.append(
                                {
                                    "text": full_desc,
                                    "metadata": {
                                        "page": slide_num,
                                        "type": "image",
                                        "source": "pptx_image",
                                        "image_index": image_count,
                                    },
                                }
                            )
                    except Exception as e:
                        pass

                if shape.has_table:
                    table = shape.table
                    for row in table.rows:
                        row_text = []
                        for cell in row.cells:
                            row_text.append(cell.text.replace("\n", " ").strip())
                        if any(row_text):
                            slide_text.append(" | ".join(row_text))

            if slide_text:
                text = "\n".join(slide_text)
                blocks.append(
                    {
                        "text": text,
                        "metadata": {
                            "page": slide_num,
                            "type": "text",
                            "source": "pptx_extraction",
                            "char_count": len(text),
                        },
                    }
                )

        return blocks
